#!/usr/bin/env python3
"""
SERESLER BİLİRKİŞİLİK — Excel çıkarım betiği (Faz 0).

Bir .xlsx dosyasını iki yoldan okur:
  1. openpyxl  → hücre formülleri + önbelleğe alınmış (cached) değerler, birleştirilmiş hücreler,
                 dizi formülleri, standart veri doğrulamaları ve koşullu biçimler.
  2. Ham XML   → x14 uzantılı veri doğrulamaları ve koşullu biçimler (openpyxl bunları atlar),
                 köprüler, çizimler/görseller ve konumları, dış bağlantılar.

Kullanım:
  python3 scripts/extract/extract_workbook.py "EĞİTİM 001.xlsx" scripts/extract/out/core
  python3 scripts/extract/extract_workbook.py "2.2 Gunes_Etkenligi_Model.xlsx" scripts/extract/out/gunes

Çıktı: <out>/_workbook.json (envanter, dış bağlantılar, görseller) ve <out>/sheets/<index>_<ad>.json
"""
import json
import os
import re
import sys
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
    "xm": "http://schemas.microsoft.com/office/excel/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
FUNC_RE = re.compile(r"\b([A-Z][A-Z0-9\.]*)\s*\(")
SHEET_REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\.ğüşıöçĞÜŞİÖÇ]+))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3})")
EXT_REF_RE = re.compile(r"\[(\d+)\]")


def safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", name)


def read_xml(z: zipfile.ZipFile, path: str):
    try:
        return ET.fromstring(z.read(path))
    except KeyError:
        return None


def rels_map(z: zipfile.ZipFile, rels_path: str):
    root = read_xml(z, rels_path)
    out = {}
    if root is None:
        return out
    for rel in root.findall("rel:Relationship", NS):
        out[rel.get("Id")] = {"target": rel.get("Target"), "type": rel.get("Type").rsplit("/", 1)[-1], "mode": rel.get("TargetMode")}
    return out


def resolve(base: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    parts = base.split("/")[:-1]
    for seg in target.split("/"):
        if seg == "..":
            parts.pop()
        elif seg and seg != ".":
            parts.append(seg)
    return "/".join(parts)


def cell_json(val):
    if isinstance(val, ArrayFormula):
        return {"array": True, "ref": val.ref, "text": val.text}
    if hasattr(val, "isoformat"):
        return val.isoformat()
    return val


def extract_sheet_xml(z, sheet_path, sheet_rels):
    """x14 veri doğrulama, x14 koşullu biçim, köprüler, çizim ilişkileri."""
    root = read_xml(z, sheet_path)
    result = {"x14DataValidations": [], "x14ConditionalFormatting": [], "dataValidations": [], "conditionalFormatting": [], "hyperlinks": [], "drawings": [], "dimension": None}
    if root is None:
        return result
    dim = root.find("m:dimension", NS)
    if dim is not None:
        result["dimension"] = dim.get("ref")

    for dv in root.findall("m:dataValidations/m:dataValidation", NS):
        f1 = dv.find("m:formula1", NS)
        f2 = dv.find("m:formula2", NS)
        result["dataValidations"].append({
            "sqref": dv.get("sqref"), "type": dv.get("type"), "operator": dv.get("operator"),
            "formula1": f1.text if f1 is not None else None, "formula2": f2.text if f2 is not None else None,
            "allowBlank": dv.get("allowBlank"), "promptTitle": dv.get("promptTitle"), "prompt": dv.get("prompt"),
        })

    for cf in root.findall("m:conditionalFormatting", NS):
        rules = []
        for rule in cf.findall("m:cfRule", NS):
            rules.append({"type": rule.get("type"), "operator": rule.get("operator"), "dxfId": rule.get("dxfId"), "priority": rule.get("priority"),
                          "formulas": [f.text for f in rule.findall("m:formula", NS)]})
        result["conditionalFormatting"].append({"sqref": cf.get("sqref"), "rules": rules})

    for ext in root.findall("m:extLst/m:ext", NS):
        for dv in ext.findall("x14:dataValidations/x14:dataValidation", NS):
            f1 = dv.find("x14:formula1/xm:f", NS)
            f2 = dv.find("x14:formula2/xm:f", NS)
            sq = dv.find("xm:sqref", NS)
            result["x14DataValidations"].append({
                "sqref": sq.text if sq is not None else None, "type": dv.get("type"), "operator": dv.get("operator"),
                "formula1": f1.text if f1 is not None else None, "formula2": f2.text if f2 is not None else None,
                "allowBlank": dv.get("allowBlank"), "promptTitle": dv.get("promptTitle"), "prompt": dv.get("prompt"),
            })
        for cf in ext.findall("x14:conditionalFormattings/x14:conditionalFormatting", NS):
            sq = cf.find("xm:sqref", NS)
            rules = []
            for rule in cf.findall("x14:cfRule", NS):
                dxf = rule.find("x14:dxf", NS)
                fill = None
                if dxf is not None:
                    for color_el in dxf.iter():
                        tag = color_el.tag.split("}")[-1]
                        if tag in ("fgColor", "bgColor", "color"):
                            fill = fill or []
                            fill.append({k: v for k, v in color_el.attrib.items()})
                    grad = dxf.find(".//m:gradientFill", NS)
                    if grad is not None:
                        fill = fill or []
                        fill.append({"gradient": grad.attrib, "stops": [{"pos": s.get("position"), **(s.find("m:color", NS).attrib if s.find("m:color", NS) is not None else {})} for s in grad.findall("m:stop", NS)]})
                rules.append({"type": rule.get("type"), "operator": rule.get("operator"), "priority": rule.get("priority"),
                              "formulas": [f.text for f in rule.findall("xm:f", NS)], "fill": fill})
            result["x14ConditionalFormatting"].append({"sqref": sq.text if sq is not None else None, "rules": rules})

    for hl in root.findall("m:hyperlinks/m:hyperlink", NS):
        rid = hl.get("{%s}id" % NS["r"])
        target = sheet_rels.get(rid, {}).get("target") if rid else None
        result["hyperlinks"].append({"ref": hl.get("ref"), "location": hl.get("location"), "display": hl.get("display"), "tooltip": hl.get("tooltip"), "externalTarget": target})

    for dr in root.findall("m:drawing", NS):
        rid = dr.get("{%s}id" % NS["r"])
        rel = sheet_rels.get(rid)
        if not rel:
            continue
        dpath = resolve(sheet_path, rel["target"])
        drels = rels_map(z, dpath.rsplit("/", 1)[0] + "/_rels/" + dpath.rsplit("/", 1)[1] + ".rels")
        droot = read_xml(z, dpath)
        if droot is None:
            continue
        for anchor in list(droot):
            tag = anchor.tag.split("}")[-1]
            frm = anchor.find("xdr:from", NS)
            to = anchor.find("xdr:to", NS)
            def pos(el):
                if el is None:
                    return None
                return {"col": int(el.find("xdr:col", NS).text), "row": int(el.find("xdr:row", NS).text) + 1,
                        "colLetter": get_column_letter(int(el.find("xdr:col", NS).text) + 1)}
            pic = anchor.find("xdr:pic", NS)
            sp = anchor.find("xdr:sp", NS)
            item = {"anchor": tag, "from": pos(frm), "to": pos(to)}
            if pic is not None:
                blip = pic.find(".//a:blip", NS)
                name_el = pic.find(".//xdr:cNvPr", NS)
                emb = blip.get("{%s}embed" % NS["r"]) if blip is not None else None
                item["kind"] = "picture"
                item["name"] = name_el.get("name") if name_el is not None else None
                item["media"] = resolve(dpath, drels[emb]["target"]) if emb in drels else None
                ext = pic.find(".//a:xfrm/a:ext", NS)
                if ext is not None and ext.get("cx") and ext.get("cy"):
                    item["sizeEmu"] = {"cx": int(ext.get("cx")), "cy": int(ext.get("cy"))}
            elif sp is not None:
                name_el = sp.find(".//xdr:cNvPr", NS)
                txt = "".join(t.text or "" for t in sp.findall(".//a:t", NS))
                item["kind"] = "shape"
                item["name"] = name_el.get("name") if name_el is not None else None
                item["text"] = txt
            else:
                item["kind"] = tag
            result["drawings"].append(item)
    return result


def extract_external_links(z):
    wb_rels = rels_map(z, "xl/_rels/workbook.xml.rels")
    links = []
    for rid, rel in wb_rels.items():
        if rel["type"] != "externalLink":
            continue
        path = resolve("xl/workbook.xml", rel["target"])
        lrels = rels_map(z, path.rsplit("/", 1)[0] + "/_rels/" + path.rsplit("/", 1)[1] + ".rels")
        root = read_xml(z, path)
        target = next((r["target"] for r in lrels.values()), None)
        sheet_names = [s.get("val") for s in root.findall(".//m:sheetName", NS)] if root is not None else []
        cached = {}
        if root is not None:
            for sd in root.findall(".//m:sheetData", NS):
                sid = int(sd.get("sheetId"))
                sname = sheet_names[sid] if sid < len(sheet_names) else str(sid)
                cells = {}
                for c in sd.findall(".//m:cell", NS):
                    v = c.find("m:v", NS)
                    cells[c.get("r")] = v.text if v is not None else None
                cached[sname] = cells
        links.append({"rId": rid, "part": path, "target": target, "sheetNames": sheet_names, "cachedValues": cached})
    return links


def main(xlsx_path: str, out_dir: str):
    os.makedirs(os.path.join(out_dir, "sheets"), exist_ok=True)
    print(f"[1/4] openpyxl formül yükü: {xlsx_path}")
    wb_f = openpyxl.load_workbook(xlsx_path, data_only=False)
    print("[2/4] openpyxl değer yükü (cached)")
    wb_v = openpyxl.load_workbook(xlsx_path, data_only=True)
    z = zipfile.ZipFile(xlsx_path)
    wb_xml = read_xml(z, "xl/workbook.xml")
    wb_rels = rels_map(z, "xl/_rels/workbook.xml.rels")
    sheet_parts = {}
    for s in wb_xml.findall("m:sheets/m:sheet", NS):
        rid = s.get("{%s}id" % NS["r"])
        sheet_parts[s.get("name")] = resolve("xl/workbook.xml", wb_rels[rid]["target"])

    inventory = []
    cross_refs = Counter()
    print("[3/4] sayfalar")
    for idx, ws in enumerate(wb_f.worksheets):
        wsv = wb_v[ws.title]
        cells = {}
        n_formula = n_value = n_array = 0
        funcs = Counter()
        refs = Counter()
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                entry = {}
                if isinstance(v, ArrayFormula):
                    n_array += 1
                    entry["f"] = v.text
                    entry["arrayRef"] = v.ref
                    entry["cached"] = cell_json(wsv[c.coordinate].value)
                    ftxt = v.text
                elif isinstance(v, str) and v.startswith("="):
                    n_formula += 1
                    entry["f"] = v
                    entry["cached"] = cell_json(wsv[c.coordinate].value)
                    ftxt = v
                else:
                    n_value += 1
                    entry["v"] = cell_json(v)
                    ftxt = None
                if ftxt:
                    for fn in FUNC_RE.findall(ftxt.upper()):
                        funcs[fn] += 1
                    for m in SHEET_REF_RE.finditer(ftxt):
                        target_sheet = m.group(1) or m.group(2)
                        refs[target_sheet] += 1
                        cross_refs[(ws.title, target_sheet)] += 1
                    for m in EXT_REF_RE.finditer(ftxt):
                        refs[f"[external {m.group(1)}]"] += 1
                        cross_refs[(ws.title, f"[external {m.group(1)}]")] += 1
                if c.hyperlink is not None:
                    entry["hyperlink"] = {"target": c.hyperlink.target, "location": c.hyperlink.location, "display": c.hyperlink.display}
                if c.comment is not None:
                    entry["comment"] = c.comment.text
                cells[c.coordinate] = entry
        part = sheet_parts.get(ws.title)
        srels = rels_map(z, part.rsplit("/", 1)[0] + "/_rels/" + part.rsplit("/", 1)[1] + ".rels") if part else {}
        xml_info = extract_sheet_xml(z, part, srels) if part else {}
        sheet_json = {
            "index": idx, "name": ws.title, "part": part, "dimension": ws.dimensions, "xmlDimension": xml_info.get("dimension"),
            "maxRow": ws.max_row, "maxCol": ws.max_column, "sheetState": ws.sheet_state,
            "counts": {"formulas": n_formula, "arrayFormulas": n_array, "values": n_value},
            "functionFrequency": dict(funcs.most_common()),
            "referencedSheets": dict(refs.most_common()),
            "mergedCells": [str(r) for r in ws.merged_cells.ranges],
            "dataValidations": xml_info.get("dataValidations", []),
            "x14DataValidations": xml_info.get("x14DataValidations", []),
            "conditionalFormatting": xml_info.get("conditionalFormatting", []),
            "x14ConditionalFormatting": xml_info.get("x14ConditionalFormatting", []),
            "hyperlinks": xml_info.get("hyperlinks", []),
            "drawings": xml_info.get("drawings", []),
            "columnWidths": {k: v.width for k, v in ws.column_dimensions.items() if v.width},
            "cells": cells,
        }
        fname = f"{idx:02d}_{safe_name(ws.title)}.json"
        with open(os.path.join(out_dir, "sheets", fname), "w", encoding="utf-8") as fh:
            json.dump(sheet_json, fh, ensure_ascii=False, indent=None)
        inventory.append({"index": idx, "name": ws.title, "file": fname, "dimension": ws.dimensions, "formulas": n_formula, "arrayFormulas": n_array, "values": n_value,
                          "x14Validations": len(sheet_json["x14DataValidations"]), "validations": len(sheet_json["dataValidations"]),
                          "x14CondFmt": len(sheet_json["x14ConditionalFormatting"]), "condFmt": len(sheet_json["conditionalFormatting"]),
                          "hyperlinks": len(sheet_json["hyperlinks"]), "drawings": len(sheet_json["drawings"]), "merged": len(sheet_json["mergedCells"]),
                          "topFunctions": dict(funcs.most_common(8)), "referencedSheets": dict(refs.most_common())})
        print(f"   {idx:2d} {ws.title:12s} {ws.dimensions:14s} f={n_formula:6d} arr={n_array} v={n_value:6d} x14dv={len(sheet_json['x14DataValidations'])} hl={len(sheet_json['hyperlinks'])} img={len(sheet_json['drawings'])}")

    print("[4/4] dış bağlantılar, görseller, tema")
    ext_links = extract_external_links(z)
    media = [n for n in z.namelist() if n.startswith("xl/media/")]
    theme = read_xml(z, "xl/theme/theme1.xml")
    theme_colors = {}
    if theme is not None:
        for el in theme.find(".//a:clrScheme", NS):
            tag = el.tag.split("}")[-1]
            child = list(el)[0]
            theme_colors[tag] = child.get("lastClr") or child.get("val")
    workbook_json = {
        "source": os.path.basename(xlsx_path), "sheets": inventory,
        "crossSheetReferences": [{"from": a, "to": b, "count": n} for (a, b), n in sorted(cross_refs.items(), key=lambda kv: (-kv[1], kv[0]))],
        "externalLinks": ext_links, "media": media, "themeColors": theme_colors,
        "definedNames": [{"name": dn.get("name"), "localSheetId": dn.get("localSheetId"), "value": dn.text} for dn in wb_xml.findall("m:definedNames/m:definedName", NS)],
    }
    with open(os.path.join(out_dir, "_workbook.json"), "w", encoding="utf-8") as fh:
        json.dump(workbook_json, fh, ensure_ascii=False, indent=2)
    media_dir = os.path.join(out_dir, "media")
    os.makedirs(media_dir, exist_ok=True)
    for m in media:
        with open(os.path.join(media_dir, os.path.basename(m)), "wb") as fh:
            fh.write(z.read(m))
    print("bitti →", out_dir)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
